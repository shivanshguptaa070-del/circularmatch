import os
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def prepare_dual_workspace_image():
    sc_dir = os.path.join(os.getcwd(), 'screenshots')
    img_sup_p = os.path.join(sc_dir, 'slide3_seller_dashboard.png')
    img_buy_p = os.path.join(sc_dir, 'slide3_buyer_dashboard.png')
    out_p = os.path.join(sc_dir, 'slide3_dual_workspaces.png')
    
    if os.path.exists(img_sup_p) and os.path.exists(img_buy_p):
        img_sup = Image.open(img_sup_p)
        img_buy = Image.open(img_buy_p)
        crop_h = 1050
        w, _ = img_sup.size
        sup_c = img_sup.crop((0, 0, w, crop_h))
        buy_c = img_buy.crop((0, 0, w, crop_h))
        gap = 40
        combined = Image.new('RGBA', (w, crop_h * 2 + gap), (6, 11, 19, 255))
        combined.paste(sup_c, (0, 0))
        combined.paste(buy_c, (0, crop_h + gap))
        combined.save(out_p)
        print("Generated dual workspace mockup for Slide 3.")

def create_base_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    BG_COLOR = RGBColor(6, 11, 19)         # #060b13
    CARD_BG = RGBColor(15, 23, 42)         # #0f172a
    BORDER_COLOR = RGBColor(30, 41, 59)    # #1e293b
    GREEN_PRIMARY = RGBColor(74, 222, 128) # #4ade80
    TEAL_ACCENT = RGBColor(45, 212, 191)   # #2dd4bf
    WHITE = RGBColor(255, 255, 255)
    MUTED = RGBColor(148, 163, 184)        # #94a3b8
    DANGER_RED = RGBColor(248, 113, 113)   # #f87171
    YELLOW_ACC = RGBColor(251, 191, 36)    # #fbbf24

    sc_dir = os.path.join(os.getcwd(), 'screenshots')

    def add_slide_base(slide, title_chip, main_title, subtitle=None):
        # 1. Background fill
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()

        # 2. Header Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.2))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = title_chip.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = GREEN_PRIMARY

        p1 = tf.add_paragraph()
        p1.text = main_title
        p1.font.size = Pt(21)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p1.space_before = Pt(4)

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.size = Pt(12.5)
            p2.font.color.rgb = TEAL_ACCENT
            p2.space_before = Pt(2)

        # 3. Footer Brand Bar
        footer = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
        ftf = footer.text_frame
        ftf.word_wrap = True
        fp = ftf.paragraphs[0]
        fp.text = "♻ CircularMatch  |  HACKDAY 1.0 Submission  |  Solo Levelling — Shivansh Gupta  |  github.com/shivanshguptaa070-del/circularmatch"
        fp.font.size = Pt(9.5)
        fp.font.color.rgb = MUTED

    # ═════════════════════════════════════════════════
    # SLIDE 1: Problem Statement
    # ═════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_base(s1, "Slide 1 of 7 — Problem Statement", 
                   "India's ₹80,000 Cr industrial waste market runs on WhatsApp, phone calls & paper receipts.",
                   "The real problem is not discovery. It is qualification.")
    
    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.8), Inches(4.3))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD_BG
    card1.line.color.rgb = BORDER_COLOR
    ctf1 = card1.text_frame
    ctf1.word_wrap = True
    ctf1.margin_left = ctf1.margin_top = ctf1.margin_right = ctf1.margin_bottom = Inches(0.25)
    
    p = ctf1.paragraphs[0]
    p.text = "🚫 30–60% Rejected at the Gate"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = DANGER_RED
    
    p = ctf1.add_paragraph()
    p.text = "Industrial buyers reject 30–60% of material upon physical arrival — wrong grade, mismatched technical specs, or undisclosed contamination causing entire truckloads to turn back at enormous freight loss."
    p.font.size = Pt(10.5)
    p.font.color.rgb = MUTED
    p.space_after = Pt(12)

    p = ctf1.add_paragraph()
    p.text = "📜 Mandatory BRSR with Zero Documentation"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = DANGER_RED

    p = ctf1.add_paragraph()
    p.text = "India's SEBI now legally mandates top 1,000 listed companies prove waste recycling in annual BRSR audits — but factories currently have zero digital documentation."
    p.font.size = Pt(10.5)
    p.font.color.rgb = MUTED
    p.space_after = Pt(12)

    p = ctf1.add_paragraph()
    p.text = "⚖️ Informal Brokers = Direct Legal Liability"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = DANGER_RED

    p = ctf1.add_paragraph()
    p.text = "Brokers move material fast but produce no compliant paper trail — creating direct regulatory liability for manufacturers filing audited disclosures with SEBI."
    p.font.size = Pt(10.5)
    p.font.color.rgb = MUTED

    bb1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.55))
    bb1.fill.solid()
    bb1.fill.fore_color.rgb = RGBColor(12, 36, 24)
    bb1.line.color.rgb = GREEN_PRIMARY
    btf1 = bb1.text_frame
    bp1 = btf1.paragraphs[0]
    bp1.text = "Bottom line: Factories have material. Buyers have requirements. Nobody can prove the material actually meets them."
    bp1.font.size = Pt(10.5)
    bp1.font.bold = True
    bp1.font.color.rgb = WHITE
    bp1.alignment = PP_ALIGN.CENTER

    img_path1 = os.path.join(sc_dir, 'slide1_landing.png')
    if os.path.exists(img_path1):
        s1.shapes.add_picture(img_path1, Inches(6.9), Inches(1.8), width=Inches(5.6))

    # ═════════════════════════════════════════════════
    # SLIDE 2: Proposed Solution
    # ═════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_base(s2, "Slide 2 of 7 — Proposed Solution",
                   "CircularMatch — AI-assisted industrial material qualification and matching platform",
                   "Standardizing plain-language waste into verified Material Passports & explainable matches.")

    features = [
        ("💬 Plain Language Input", "Supplier describes waste naturally — Gemini AI structures it instantly into a standardized, schema-validated record."),
        ("📋 Material Passport", "Every material gets a passport tracking grade, quantity, lot location, evidence state, and eligibility gates."),
        ("🎯 6-Dimension Engine", "Proprietary matching algorithm scores lots against exact buyer acceptance criteria deterministically."),
        ("💡 Explainable AI Reasoning", "Gemini generates readable 'Why this match?' justifications detailing strengths, risks, and recommendations.")
    ]

    for i, (f_title, f_desc) in enumerate(features):
        x = Inches(0.8 + (i % 2) * 2.95)
        y = Inches(1.8 + (i // 2) * 1.45)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.85), Inches(1.35))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = Inches(0.12)
        cp1 = ctf.paragraphs[0]
        cp1.text = f_title
        cp1.font.size = Pt(11.5)
        cp1.font.bold = True
        cp1.font.color.rgb = GREEN_PRIMARY
        cp2 = ctf.add_paragraph()
        cp2.text = f_desc
        cp2.font.size = Pt(9.5)
        cp2.font.color.rgb = MUTED
        cp2.space_before = Pt(3)

    flow_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.85), Inches(5.8), Inches(1.95))
    flow_box.fill.solid()
    flow_box.fill.fore_color.rgb = CARD_BG
    flow_box.line.color.rgb = GREEN_PRIMARY
    ftf = flow_box.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = Inches(0.15)
    fp = ftf.paragraphs[0]
    fp.text = "THE CORE PLATFORM FLOW"
    fp.font.size = Pt(10)
    fp.font.bold = True
    fp.font.color.rgb = TEAL_ACCENT
    
    flow_steps = [
        "1. Plain language input (generator speaks or types naturally)",
        "2. AI extraction → structured JSON material record",
        "3. Material Passport with digital evidence states",
        "4. 6-factor deterministic matching against buyer specs",
        "5. Ranked results with explainable AI reasoning",
        "6. Sample request → offer → shipment chain-of-custody"
    ]
    for step in flow_steps:
        sp = ftf.add_paragraph()
        sp.text = step
        sp.font.size = Pt(9.5)
        sp.font.color.rgb = WHITE
        sp.space_before = Pt(2)

    img_path2 = os.path.join(sc_dir, 'slide2_passport.png')
    if os.path.exists(img_path2):
        s2.shapes.add_picture(img_path2, Inches(6.9), Inches(1.8), width=Inches(5.6))

    # ═════════════════════════════════════════════════
    # SLIDE 3: Target Users (Fixed Layout & Dual Mockup)
    # ═════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_base(s3, "Slide 3 of 7 — Target Users",
                   "Two sides. One regulatory & financial imperative.",
                   "Solving compliance failure for sellers and gate rejection financial loss for buyers.")

    # Left Top: Supplier Side Card
    sup_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.8), Inches(2.1))
    sup_card.fill.solid()
    sup_card.fill.fore_color.rgb = CARD_BG
    sup_card.line.color.rgb = GREEN_PRIMARY
    s_tf = sup_card.text_frame
    s_tf.word_wrap = True
    s_tf.margin_left = s_tf.margin_top = s_tf.margin_right = s_tf.margin_bottom = Inches(0.18)
    sp = s_tf.paragraphs[0]
    sp.text = "🏭 SUPPLIER SIDE: EHS / Operations Manager"
    sp.font.size = Pt(12.5)
    sp.font.bold = True
    sp.font.color.rgb = GREEN_PRIMARY

    sp = s_tf.add_paragraph()
    sp.text = "Profile: BSE/NSE-listed manufacturing enterprise in Noida/NCR generating recurring industrial waste (PET, cotton, corrugated) with mandatory SEBI BRSR Core filing."
    sp.font.size = Pt(9.5)
    sp.font.color.rgb = WHITE
    sp.space_before = Pt(4)

    sp = s_tf.add_paragraph()
    sp.text = "Pain Point: BRSR auditors reject informal broker receipts. They need verifiable digital chain-of-custody for statutory audits — currently zero digital records exist."
    sp.font.size = Pt(9.5)
    sp.font.color.rgb = DANGER_RED
    sp.space_before = Pt(4)

    # Left Bottom: Buyer Side Card
    buy_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.05), Inches(5.8), Inches(2.05))
    buy_card.fill.solid()
    buy_card.fill.fore_color.rgb = CARD_BG
    buy_card.line.color.rgb = TEAL_ACCENT
    b_tf = buy_card.text_frame
    b_tf.word_wrap = True
    b_tf.margin_left = b_tf.margin_top = b_tf.margin_right = b_tf.margin_bottom = Inches(0.18)
    bp = b_tf.paragraphs[0]
    bp.text = "🔄 BUYER SIDE: Procurement Manager"
    bp.font.size = Pt(12.5)
    bp.font.bold = True
    bp.font.color.rgb = TEAL_ACCENT

    bp = b_tf.add_paragraph()
    bp.text = "Profile: Sourcing manager at registered recyclers or brand owners building EPR compliance supply chains — needing pre-qualified, spec-matched material lots."
    bp.font.size = Pt(9.5)
    bp.font.color.rgb = WHITE
    bp.space_before = Pt(4)

    bp = b_tf.add_paragraph()
    bp.text = "Pain Point: 30–60% gate rejection rate on unverified suppliers. Every rejected shipment loses 100% freight and logistics cost with zero usable material received."
    bp.font.size = Pt(9.5)
    bp.font.color.rgb = DANGER_RED
    bp.space_before = Pt(4)

    # Bottom Banner for Slide 3
    bb3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.55))
    bb3.fill.solid()
    bb3.fill.fore_color.rgb = RGBColor(12, 36, 24)
    bb3.line.color.rgb = GREEN_PRIMARY
    bbtf3 = bb3.text_frame
    bbp3 = bbtf3.paragraphs[0]
    bbp3.text = "Why these two: The supplier has a regulatory mandate making documentation non-optional. The buyer has a financial loss making pre-qualification worth paying for."
    bbp3.font.size = Pt(10)
    bbp3.font.bold = True
    bbp3.font.color.rgb = WHITE
    bbp3.alignment = PP_ALIGN.CENTER

    # Right: Dual Workspace Image (Seller & Buyer Dashboards)
    img_dual = os.path.join(sc_dir, 'slide3_dual_workspaces.png')
    if os.path.exists(img_dual):
        s3.shapes.add_picture(img_dual, Inches(6.9), Inches(1.8), width=Inches(5.6))

    # ═════════════════════════════════════════════════
    # SLIDE 4: Technical Approach
    # ═════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_base(s4, "Slide 4 of 7 — Technical Approach",
                   "Production Stack with Explainable AI & Deterministic Matching Engine",
                   "FastAPI + React 18 + Supabase PostgreSQL + Google Gemini 2.0 Flash with Rule-based fallback.")

    rows = [
        ("Frontend", "React 18 + TypeScript + Vite + Tailwind CSS"),
        ("Backend", "FastAPI (Python) + Uvicorn + Pydantic v2"),
        ("Database", "Supabase PostgreSQL (migrations written)"),
        ("AI Core", "Google Gemini 2.0 Flash (structured JSON output)"),
        ("AI Fallback", "Deterministic rule-based regex extractor (zero API failure)"),
        ("Maps", "Leaflet — Delhi NCR match routes rendered"),
        ("Deploy", "Vercel (Frontend) + Render (API backend)")
    ]
    
    t_shape = s4.shapes.add_table(len(rows) + 1, 2, Inches(0.8), Inches(1.8), Inches(5.8), Inches(2.2))
    table = t_shape.table
    table.columns[0].width = Inches(1.4)
    table.columns[1].width = Inches(4.4)

    cell_hdr0 = table.cell(0, 0)
    cell_hdr0.text = "LAYER"
    cell_hdr0.fill.solid()
    cell_hdr0.fill.fore_color.rgb = CARD_BG
    cell_hdr0.text_frame.paragraphs[0].font.bold = True
    cell_hdr0.text_frame.paragraphs[0].font.size = Pt(9.5)
    cell_hdr0.text_frame.paragraphs[0].font.color.rgb = GREEN_PRIMARY

    cell_hdr1 = table.cell(0, 1)
    cell_hdr1.text = "TECHNOLOGY"
    cell_hdr1.fill.solid()
    cell_hdr1.fill.fore_color.rgb = CARD_BG
    cell_hdr1.text_frame.paragraphs[0].font.bold = True
    cell_hdr1.text_frame.paragraphs[0].font.size = Pt(9.5)
    cell_hdr1.text_frame.paragraphs[0].font.color.rgb = GREEN_PRIMARY

    for row_idx, (layer, tech) in enumerate(rows, start=1):
        c0 = table.cell(row_idx, 0)
        c0.text = layer
        c0.fill.solid()
        c0.fill.fore_color.rgb = CARD_BG
        c0.text_frame.paragraphs[0].font.size = Pt(9)
        c0.text_frame.paragraphs[0].font.bold = True
        c0.text_frame.paragraphs[0].font.color.rgb = TEAL_ACCENT

        c1 = table.cell(row_idx, 1)
        c1.text = tech
        c1.fill.solid()
        c1.fill.fore_color.rgb = CARD_BG
        c1.text_frame.paragraphs[0].font.size = Pt(9)
        c1.text_frame.paragraphs[0].font.color.rgb = WHITE

    feat_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(5.8), Inches(2.55))
    feat_card.fill.solid()
    feat_card.fill.fore_color.rgb = CARD_BG
    feat_card.line.color.rgb = BORDER_COLOR
    ftf4 = feat_card.text_frame
    ftf4.word_wrap = True
    ftf4.margin_left = ftf4.margin_top = ftf4.margin_right = ftf4.margin_bottom = Inches(0.15)
    
    fp = ftf4.paragraphs[0]
    fp.text = "TWO GENUINE AI FEATURES"
    fp.font.size = Pt(11)
    fp.font.bold = True
    fp.font.color.rgb = GREEN_PRIMARY

    fp = ftf4.add_paragraph()
    fp.text = "Feature 1 — Natural Language Listing Extraction:"
    fp.font.size = Pt(10)
    fp.font.bold = True
    fp.font.color.rgb = WHITE
    fp.space_before = Pt(4)

    fp = ftf4.add_paragraph()
    fp.text = "Supplier types: '2 tonnes cotton cutting waste from Noida, Grade A, monthly' → Gemini extracts material_type, qty, unit, grade, location, contamination level instantly."
    fp.font.size = Pt(9)
    fp.font.color.rgb = MUTED

    fp = ftf4.add_paragraph()
    fp.text = "Feature 2 — Explainable Match Reasoning ('Why this match?'):"
    fp.font.size = Pt(10)
    fp.font.bold = True
    fp.font.color.rgb = WHITE
    fp.space_before = Pt(4)

    fp = ftf4.add_paragraph()
    fp.text = "Deterministic 6-D engine calculates score → Gemini generates clear 2-3 sentence explanation of strengths and risks. (AI explains the score — does NOT determine it)."
    fp.font.size = Pt(9)
    fp.font.color.rgb = MUTED

    img_path4 = os.path.join(sc_dir, 'slide4_match_detail.png')
    if os.path.exists(img_path4):
        s4.shapes.add_picture(img_path4, Inches(6.9), Inches(1.8), width=Inches(5.6))

    # ═════════════════════════════════════════════════
    # SLIDE 5: Market & Business Potential
    # ═════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_base(s5, "Slide 5 of 7 — Market & Business Potential",
                   "Monetizing the Compliance Trail, Not Just Transactions",
                   "SEBI BRSR Core + EPR Mandate create an un-bypassable SaaS compliance lock-in.")

    mkt_metrics = [
        ("TAM — India Industrial Secondary Materials", "₹80,000–1,20,000 Cr", "Annual industrial recyclable flows across manufacturing sectors"),
        ("SAM — NCR Industrial Clusters", "₹3,000–10,000 Cr", "Delhi NCR manufacturing hubs (Noida, Gr. Noida, Faridabad, Gurugram)"),
        ("Immediate Market Wedge", "500 Listed Companies", "Mandatory BRSR Core reporting companies this fiscal year → 1,000 by FY27")
    ]

    for idx, (title, val, desc) in enumerate(mkt_metrics):
        box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8 + idx * 0.95), Inches(5.8), Inches(0.85))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = BORDER_COLOR
        btf = box.text_frame
        btf.word_wrap = True
        btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = Inches(0.1)
        bp = btf.paragraphs[0]
        bp.text = f"{title.upper()}:  {val}"
        bp.font.size = Pt(11)
        bp.font.bold = True
        bp.font.color.rgb = GREEN_PRIMARY if idx == 0 else TEAL_ACCENT if idx == 1 else YELLOW_ACC
        bp2 = btf.add_paragraph()
        bp2.text = desc
        bp2.font.size = Pt(9.5)
        bp2.font.color.rgb = MUTED

    biz_rows = [
        ("Hook", "Free listing + matching", "₹0", "Acquires both sides rapidly"),
        ("Revenue", "EHS Compliance SaaS", "₹5K–15K/mo", "BRSR audit trail lives here — bypass = failure"),
        ("Upsell", "Verification per lot", "₹500–5K/lot", "Eliminates buyer gate rejection risk")
    ]
    b_table_shape = s5.shapes.add_table(len(biz_rows) + 1, 4, Inches(0.8), Inches(4.75), Inches(5.8), Inches(1.3))
    b_table = b_table_shape.table
    b_table.columns[0].width = Inches(0.9)
    b_table.columns[1].width = Inches(1.8)
    b_table.columns[2].width = Inches(1.1)
    b_table.columns[3].width = Inches(2.0)

    for c_i, h in enumerate(["LAYER", "PRODUCT", "PRICE", "WHY STICKY"]):
        cell = b_table.cell(0, c_i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(9)
        cell.text_frame.paragraphs[0].font.color.rgb = GREEN_PRIMARY

    for r_i, r_data in enumerate(biz_rows, start=1):
        for c_i, val in enumerate(r_data):
            cell = b_table.cell(r_i, c_i)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            cell.text_frame.paragraphs[0].font.size = Pt(8.5)
            cell.text_frame.paragraphs[0].font.color.rgb = WHITE

    bb5 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(5.8), Inches(0.65))
    bb5.fill.solid()
    bb5.fill.fore_color.rgb = RGBColor(12, 36, 24)
    bb5.line.color.rgb = GREEN_PRIMARY
    bbtf5 = bb5.text_frame
    bbtf5.word_wrap = True
    bbtf5.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = Inches(0.08)
    bbp5 = bbtf5.paragraphs[0]
    bbp5.text = "The Un-bypassable Insight: We don't charge transaction fees. We charge for the compliance audit trail. Leave platform = lose audit record."
    bbp5.font.size = Pt(9)
    bbp5.font.bold = True
    bbp5.font.color.rgb = WHITE

    img_path5 = os.path.join(sc_dir, 'slide5_impact.png')
    if os.path.exists(img_path5):
        s5.shapes.add_picture(img_path5, Inches(6.9), Inches(1.8), width=Inches(5.6))

    # ═════════════════════════════════════════════════
    # SLIDE 6: Scalability & Future
    # ═════════════════════════════════════════════════
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_base(s6, "Slide 6 of 7 — Scalability & Future",
                   "Modular Architecture Built to Scale from NCR to National Coverage",
                   "Production-ready design with configurable weights, pluggable AI, and multi-tenant schema.")

    arch_items = [
        ("🗄️ In-Memory to Supabase Swap", "Zero code changes to go live; database migrations and Supabase Auth schemas already written."),
        ("🤖 Pluggable AI Service Layer", "Swap Gemini model, temperature, or provider without touching business matching logic."),
        ("⚖️ Configurable Weights Engine", "Admins adjust dimension weights and eligibility gate thresholds directly via built-in dashboard."),
        ("🏢 Multi-Company RBAC Schema", "Multi-tenant organization hierarchy and role permissions already structured in Postgres.")
    ]
    for i, (a_t, a_d) in enumerate(arch_items):
        x = Inches(0.8 + (i % 2) * 2.95)
        y = Inches(1.8 + (i // 2) * 1.35)
        box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.85), Inches(1.25))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = BORDER_COLOR
        btf = box.text_frame
        btf.word_wrap = True
        btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = Inches(0.1)
        bp = btf.paragraphs[0]
        bp.text = a_t
        bp.font.size = Pt(10.5)
        bp.font.bold = True
        bp.font.color.rgb = GREEN_PRIMARY
        bp2 = btf.add_paragraph()
        bp2.text = a_d
        bp2.font.size = Pt(9)
        bp2.font.color.rgb = MUTED
        bp2.space_before = Pt(2)

    rm_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.6), Inches(5.8), Inches(1.4))
    rm_box.fill.solid()
    rm_box.fill.fore_color.rgb = CARD_BG
    rm_box.line.color.rgb = TEAL_ACCENT
    rtf = rm_box.text_frame
    rtf.word_wrap = True
    rtf.margin_left = rtf.margin_top = rtf.margin_right = rtf.margin_bottom = Inches(0.12)
    rp = rtf.paragraphs[0]
    rp.text = "12-MONTH COMMERCIAL ROADMAP"
    rp.font.size = Pt(10)
    rp.font.bold = True
    rp.font.color.rgb = TEAL_ACCENT

    roadmap_steps = [
        "Month 1–3:  5 paying EHS pilot customers in Noida industrial zone",
        "Month 3–6:  First 100 live transactions with digital chain-of-custody",
        "Month 6–9:  Automated BRSR certificate PDF generation live (NCR expansion)",
        "Month 9–12: Pre-seed funding closed (Greenovation target: Oct 31, 2026)"
    ]
    for step in roadmap_steps:
        sp = rtf.add_paragraph()
        sp.text = step
        sp.font.size = Pt(9)
        sp.font.color.rgb = WHITE
        sp.space_before = Pt(2)

    bb6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.1), Inches(5.8), Inches(0.7))
    bb6.fill.solid()
    bb6.fill.fore_color.rgb = RGBColor(12, 36, 24)
    bb6.line.color.rgb = GREEN_PRIMARY
    bbtf6 = bb6.text_frame
    bbtf6.word_wrap = True
    bbtf6.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = Inches(0.08)
    bbp6 = bbtf6.paragraphs[0]
    bbp6.text = "Data Moat: First mover with labeled secondary material quality outcomes and rejection reasons owns the pricing intelligence layer across India."
    bbp6.font.size = Pt(9)
    bbp6.font.bold = True
    bbp6.font.color.rgb = WHITE

    img_path6 = os.path.join(sc_dir, 'slide6_admin.png')
    if os.path.exists(img_path6):
        s6.shapes.add_picture(img_path6, Inches(6.9), Inches(1.8), width=Inches(5.6))

    # ═════════════════════════════════════════════════
    # SLIDE 7: If We Had More Time
    # ═════════════════════════════════════════════════
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_base(s7, "Slide 7 of 7 — If We Had More Time",
                   "Roadmap to Mission-Critical Industrial Compliance Infrastructure",
                   "Transforming an AI matching MVP into India's certified secondary material registry.")

    priorities = [
        ("1. BRSR Compliance Certificate Generator", "Auto-generate cryptographically signed chain-of-custody PDF directly usable in SEBI audit submissions. Converts CircularMatch into compliance infrastructure."),
        ("2. AI Photo Quality Grading (Computer Vision)", "Supplier uploads scrap photo → computer vision returns contamination score and material grade before shipment. Eliminates the 30–60% gate rejection rate."),
        ("3. Live Supabase Connection", "Database migrations are fully written; Supabase Auth schema exists. One environment variable away from persistent production cloud data."),
        ("4. EPR Credit Chain Documentation", "Connect EPR-obligated corporate brands directly with CPCB-registered recyclers with traceable plastic credits and verified disposal trails."),
        ("5. EU Digital Product Passport (DPP)", "JSON-LD material passports with GS1 QR codes. Indian exporters need ESPR compliance by Feb 2027. Passport architecture maps directly to this standard.")
    ]

    for idx, (p_title, p_desc) in enumerate(priorities):
        box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8 + idx * 0.85), Inches(5.8), Inches(0.78))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = BORDER_COLOR
        btf = box.text_frame
        btf.word_wrap = True
        btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = Inches(0.08)
        bp = btf.paragraphs[0]
        bp.text = p_title
        bp.font.size = Pt(10.5)
        bp.font.bold = True
        bp.font.color.rgb = GREEN_PRIMARY
        bp2 = btf.add_paragraph()
        bp2.text = p_desc
        bp2.font.size = Pt(8.5)
        bp2.font.color.rgb = MUTED

    sub_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(5.8), Inches(0.65))
    sub_box.fill.solid()
    sub_box.fill.fore_color.rgb = RGBColor(15, 23, 42)
    sub_box.line.color.rgb = GREEN_PRIMARY
    stf = sub_box.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sp.text = "SUBMISSION: CircularMatch | Team: Solo Levelling (Shivansh Gupta) | GitHub: github.com/shivanshguptaa070-del/circularmatch"
    sp.font.size = Pt(9)
    sp.font.bold = True
    sp.font.color.rgb = WHITE
    sp.alignment = PP_ALIGN.CENTER

    img_path7 = os.path.join(sc_dir, 'slide7_map.png')
    if os.path.exists(img_path7):
        s7.shapes.add_picture(img_path7, Inches(6.9), Inches(1.8), width=Inches(5.6))

    out_file = os.path.join(os.getcwd(), "CircularMatch_HACKDAY_Pitch.pptx")
    prs.save(out_file)
    print(f"Base PowerPoint presentation saved to: {out_file}")

if __name__ == '__main__':
    prepare_dual_workspace_image()
    create_base_presentation()
